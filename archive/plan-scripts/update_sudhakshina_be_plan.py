import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import copy

stories = [
    # --- PHASE 1: Zero-Dependency Foundations & Interactive Prototype (Weeks 1 - 4: Oct 12 - Nov 8) ---
    {
        "epic": "Foundations & Auth",
        "feature": "Dev Environment Setup",
        "story": "Monorepo and App Scaffolding: React 19 + Vite Frontend, FastAPI Backend, Pydantic and Linters",
        "layer": "BE",
        "owner": "Vishwachandra",
        "weeks": ["W1"]
    },
    {
        "epic": "Foundations & Auth",
        "feature": "UI Framework & Base Layout",
        "story": "Setup Base Application Styling, Common UI Components and Page Layout",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W1", "W2"]
    },
    {
        "epic": "Foundations & Auth",
        "feature": "Login Page UI",
        "story": "Login Page UI: Welcome card, Mars-Federated-SSO badge, email input (effem.com) and Sign-In button",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W1", "W2"]
    },
    {
        "epic": "Foundations & Auth",
        "feature": "API Contracts & Mock Services",
        "story": "Define API Data Models and Setup Mock APIs for Frontend Development",
        "layer": "BE",
        "owner": "Sudhakshina",
        "weeks": ["W1", "W2"]
    },
    {
        "epic": "Foundations & Auth",
        "feature": "Federated SSO & Auth API",
        "story": "Azure AD / Entra ID JWT Validation Middleware and Security Controls",
        "layer": "BE",
        "owner": "Sudhakshina",
        "weeks": ["W2", "W3"]
    },
    {
        "epic": "Main Page & Parameters",
        "feature": "KBI Header Bar & Tabs UI",
        "story": "Main Page KBI Header Bar (RSV, Share, Volume, Pen.) with CMI / Finance / Brand Category Tabs",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W2", "W3"]
    },
    {
        "epic": "Main Page & Parameters",
        "feature": "Parameter Selection Form UI",
        "story": "Parameter Selection Form: Region, Market, Presentation Template dropdowns and client validation",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W3", "W4 - M1"]
    },
    {
        "epic": "Main Page & Parameters",
        "feature": "Recent FNOS Reports UI",
        "story": "Recent FNOS Reports List Cards (FNOS_26_Q1/Q2/Q3, View: Total Europe Cluster, Status Badges)",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W3", "W4 - M1"]
    },
    {
        "epic": "Presentation Canvas & Slides",
        "feature": "Slide Thumbnail Navigation",
        "story": "Left Sidebar Slide Thumbnail Navigation: Mini slide preview list and active slide selection",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W3", "W4 - M1"]
    },
    {
        "epic": "Presentation Canvas & Slides",
        "feature": "16:9 Slide Canvas Viewer",
        "story": "Main Slide Canvas: 16:9 Slide Viewer and Interactive Data Display",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W4 - M1"]
    },
    {
        "epic": "Distribution & PPTX Engine",
        "feature": "PowerPoint Template Setup",
        "story": "PowerPoint Template Engine: Deconstruct Master Deck into Reusable Slide Layouts",
        "layer": "BE",
        "owner": "Sudhakshina",
        "weeks": ["W4 - M1"]
    },

    # --- PHASE 2: Live Data Integration & Orchestration (Weeks 5 - 8: Nov 9 - Dec 6) ---
    {
        "epic": "Main Page & Parameters",
        "feature": "Database Connection & ORM",
        "story": "Azure SQL DB Connection and SQLAlchemy ORM Layer (Consuming DS-managed schema and application tables)",
        "layer": "BE",
        "owner": "Vishwachandra",
        "weeks": ["W5", "W6"]
    },
    {
        "epic": "Main Page & Parameters",
        "feature": "Parameter & Report APIs",
        "story": "REST Endpoints for KBI metrics, filter options and recent report history (Consuming live DS tables)",
        "layer": "BE",
        "owner": "Sudhakshina",
        "weeks": ["W5", "W6"]
    },
    {
        "epic": "Presentation Canvas & Slides",
        "feature": "Analytics Data Integration",
        "story": "Connect Backend to Databricks to Fetch Brand Performance Metrics",
        "layer": "BE",
        "owner": "Sudhakshina",
        "weeks": ["W5", "W6"]
    },
    {
        "epic": "AI Co-Pilot & Commentary",
        "feature": "Slide Commentary Editor",
        "story": "Slide Commentary Editor: In-Canvas Narrative Editing and Save-to-Slide UI",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W6", "W7"]
    },
    {
        "epic": "Presentation Canvas & Slides",
        "feature": "Data Caching & Performance",
        "story": "Implement Data Caching for Fast Slide Loading and Navigation",
        "layer": "BE",
        "owner": "Sudhakshina",
        "weeks": ["W6", "W7"]
    },
    {
        "epic": "Presentation Canvas & Slides",
        "feature": "Presentation Orchestrator Engine",
        "story": "Setup Background Job Processing to Manage Long-Running Deck Generation",
        "layer": "BE",
        "owner": "Vishwachandra",
        "weeks": ["W6", "W7"]
    },
    {
        "epic": "Presentation Canvas & Slides",
        "feature": "Real-Time Progress Streaming",
        "story": "Report Generation Job Progress Tracking via Server-Sent Events (SSE)",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W7", "W8"]
    },
    {
        "epic": "AI Co-Pilot & Commentary",
        "feature": "Right Side AI & Commentary Panel",
        "story": "Right Side Panel: AI Insights Assistant and Slide Commentary Editor",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W7", "W8"]
    },
    {
        "epic": "Presentation Canvas & Slides",
        "feature": "Deck Slide Synchronization",
        "story": "Deck Navigation: Active Slide Switching and Data Synchronization across 80 Slides",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W7", "W8"]
    },
    {
        "epic": "AI Co-Pilot & Commentary",
        "feature": "Slide Edit Persistence API",
        "story": "Slide Narrative Delta Persistence API (Storing user edits in DS-provided Azure SQL tables)",
        "layer": "BE",
        "owner": "Sudhakshina",
        "weeks": ["W8"]
    },

    # --- PHASE 3: GenAI Synthesis, PPTX Engine & Distribution (Weeks 9 - 12: Dec 7 - Jan 3) ---
    {
        "epic": "Distribution & PPTX Engine",
        "feature": "Distribution Modal Dialog UI",
        "story": "Distribute Report Modal Dialog: Title input, recipient group pills (Executive Board/Managers) and Summary",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W9"]
    },
    {
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Azure OpenAI Integration",
        "story": "Setup Secure Azure OpenAI Service to Generate Structured Slide Commentary",
        "layer": "BE",
        "owner": "Sudhakshina",
        "weeks": ["W9", "W10"]
    },
    {
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Slide Narrative Generation",
        "story": "Build AI Commentary Generation API with Business Rules and Guardrails",
        "layer": "BE",
        "owner": "Sudhakshina",
        "weeks": ["W9", "W10"]
    },
    {
        "epic": "Distribution & PPTX Engine",
        "feature": "Executive Summary & Regeneration",
        "story": "Generate AI Executive Summary with Live Regeneration Action in Distribution Modal",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W9", "W10"]
    },
    {
        "epic": "Distribution & PPTX Engine",
        "feature": "Enterprise Email Service",
        "story": "Enterprise Email Integration (Microsoft Graph API / SendGrid): Dispatch Report to Selected Groups",
        "layer": "BE",
        "owner": "Sudhakshina",
        "weeks": ["W10", "W11"]
    },
    {
        "epic": "Distribution & PPTX Engine",
        "feature": "Automated PowerPoint Generator",
        "story": "Build Serverless Worker to Populate and Generate Full 80-Slide PowerPoint Deck (Azure function)",
        "layer": "BE",
        "owner": "Vishwachandra",
        "weeks": ["W10", "W11"]
    },
    {
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Interactive AI Assistant",
        "story": "Implement Conversational Q&A Assistant Streaming Backend Endpoint in Right Panel",
        "layer": "BE",
        "owner": "Sudhakshina",
        "weeks": ["W10", "W11"]
    },
    {
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Slide Chart Components",
        "story": "Build Interactive Chart Visualizations for Brand Share and Trend Slides",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W11", "W12"]
    },
    {
        "epic": "Distribution & PPTX Engine",
        "feature": "Secure Report Download",
        "story": "Generate Secure Temporary Download Links to Export PowerPoint Files from Azure Storage",
        "layer": "BE",
        "owner": "Sudhakshina",
        "weeks": ["W11", "W12"]
    },
    {
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Accessibility & Polish",
        "story": "Accessibility, Cross-Browser Compatibility and UI Polish",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W12"]
    },

    # --- PHASE 4: Hardening, Latency Profiling & Final Code Freeze (Weeks 13 - 14: Jan 4 - Jan 17) ---
    {
        "epic": "Hardening & Code Freeze",
        "feature": "Frontend State Optimization",
        "story": "Frontend Rendering Optimization, Memory Leak Prevention, Error Boundaries and Toast System",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W13"]
    },
    {
        "epic": "Hardening & Code Freeze",
        "feature": "Backend Query Profiling & Tuning",
        "story": "Analytical Query Profiling, In-Memory Caching Rules, Connection Pooling and Async Thread Tuning",
        "layer": "BE",
        "owner": "Sudhakshina",
        "weeks": ["W13"]
    },
    {
        "epic": "Hardening & Code Freeze",
        "feature": "80-Slide Rendering Latency Tuning",
        "story": "80-Slide PowerPoint Generation Optimization (Sub-second target) and Memory Profiling",
        "layer": "BE",
        "owner": "Vishwachandra",
        "weeks": ["W13", "W14"]
    },
    {
        "epic": "Hardening & Code Freeze",
        "feature": "Final Feature Freeze & Release Build",
        "story": "Final Production Backend Packaging, Container Build and Handoff to Testing Team",
        "layer": "BE",
        "owner": "Sudhakshina",
        "weeks": ["W14"]
    }
]

week_cols = [
    ("W1", "W1"),
    ("W2", "W2"),
    ("W3", "W3"),
    ("W4 - M1", "W4 - M1"),
    ("W5", "W5"),
    ("W6", "W6"),
    ("W7", "W7"),
    ("W8", "W8"),
    ("W9", "W9"),
    ("W10", "W10"),
    ("W11", "W11"),
    ("W12", "W12"),
    ("W13", "W13"),
    ("W14", "W14")
]

headers = ["Epic", "Feature", "User Story", "Layer", "Owner"] + [w[1] for w in week_cols]

def update_workbook(filepath, ref_ws):
    wb = openpyxl.load_workbook(filepath)
    if "Webapp Plan" in wb.sheetnames:
        del wb["Webapp Plan"]
    ws = wb.create_sheet(title="Webapp Plan")

    font_header = copy.copy(ref_ws.cell(1, 1).font)
    font_data = copy.copy(ref_ws.cell(2, 1).font)
    border_cell = copy.copy(ref_ws.cell(1, 1).border)
    fill_empty = PatternFill(fill_type=None)
    fill_m1 = copy.copy(ref_ws.cell(1, 8).fill)
    fill_active = copy.copy(ref_ws.cell(2, 5).fill)

    ws.append(headers)
    ws.row_dimensions[1].height = 14.25

    for c_idx in range(1, len(headers) + 1):
        cell = ws.cell(1, c_idx)
        cell.font = font_header
        cell.border = border_cell
        if headers[c_idx - 1] == "W4 - M1":
            cell.fill = fill_m1
        else:
            cell.fill = fill_empty
        if c_idx >= 6:
            cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx, s in enumerate(stories, start=2):
        row_vals = [
            s["epic"],
            s["feature"],
            s["story"],
            s["layer"],
            s["owner"]
        ]
        for _ in week_cols:
            row_vals.append(None)
            
        ws.append(row_vals)
        ws.row_dimensions[r_idx].height = 14.25

        for c_idx in range(1, 6):
            cell = ws.cell(r_idx, c_idx)
            cell.font = font_data
            cell.border = border_cell
            cell.fill = fill_empty
            if c_idx in (2, 3):
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            elif c_idx in (4, 5):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(vertical="center")

        for w_offset, (w_code, _) in enumerate(week_cols):
            col_pos = 6 + w_offset
            w_cell = ws.cell(r_idx, col_pos)
            w_cell.border = border_cell
            if w_code in s["weeks"]:
                w_cell.fill = fill_active
            else:
                w_cell.fill = fill_empty

    col_widths = {
        "A": 22.0,
        "B": 28.38,
        "C": 48.0,
        "D": 10.0,
        "E": 20.38,
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width
    for w_col_idx in range(6, len(headers) + 1):
        col_letter = get_column_letter(w_col_idx)
        if headers[w_col_idx - 1] == "W4 - M1":
            ws.column_dimensions[col_letter].width = 6.13
        else:
            ws.column_dimensions[col_letter].width = 5.0

    wb.save(filepath)
    print(f"Successfully updated {filepath}!")

if __name__ == "__main__":
    ref_wb = openpyxl.load_workbook("FNOS_Project_Plan_Initial_Draft.xlsx")
    ref_ws = ref_wb["DS Team Plan"]

    # 1. Update FNOS_App_Development_Plan.xlsx
    update_workbook("FNOS_App_Development_Plan.xlsx", ref_ws)

    # 2. Update FNOS_Project_Plan_Initial_Draft.xlsx
    update_workbook("FNOS_Project_Plan_Initial_Draft.xlsx", ref_ws)
