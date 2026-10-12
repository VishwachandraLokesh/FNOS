import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

stories = [
    # --- PHASE 1: Zero-Dependency Foundations & Interactive Prototype (Weeks 1 - 4: Oct 12 - Nov 8) ---
    {
        "epic": "Foundations & Auth",
        "feature": "Dev Environment Setup",
        "story": "Monorepo and App Scaffolding: React 19 + Vite Frontend, FastAPI Backend, Pydantic and Linters",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "12-Oct-2026",
        "end": "18-Oct-2026",
        "weeks": ["W1"]
    },
    {
        "epic": "Foundations & Auth",
        "feature": "Design System & UI Tokens",
        "story": "Mars Design System Setup: UI tokens, typography, 16:9 canvas container shell and base styling",
        "layer": "FE",
        "owner": "Sudhakshina",
        "start": "12-Oct-2026",
        "end": "25-Oct-2026",
        "weeks": ["W1", "W2"]
    },
    {
        "epic": "Foundations & Auth",
        "feature": "Login Page UI",
        "story": "Login Page UI: Welcome card, Mars-Federated-SSO badge, email input (effem.com) and Sign-In button",
        "layer": "FE",
        "owner": "Sudhakshina",
        "start": "12-Oct-2026",
        "end": "25-Oct-2026",
        "weeks": ["W1", "W2"]
    },
    {
        "epic": "Foundations & Auth",
        "feature": "API Contracts & Stubs",
        "story": "OpenAPI and Pydantic Data Contract Specifications and Mock Fixture Generator for Local Dev",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "12-Oct-2026",
        "end": "25-Oct-2026",
        "weeks": ["W1", "W2"]
    },
    {
        "epic": "Foundations & Auth",
        "feature": "Federated SSO & Auth API",
        "story": "Azure AD / Entra ID Federated SSO Integration: MSAL.js Auth Context (FE) and JWT Validation Middleware",
        "layer": "FE + BE",
        "owner": "Sudhakshina",
        "start": "19-Oct-2026",
        "end": "01-Nov-2026",
        "weeks": ["W2", "W3"]
    },
    {
        "epic": "Main Page & Parameters",
        "feature": "KBI Header Bar & Tabs UI",
        "story": "Main Page KBI Header Bar (RSV, Share, Volume, Pen.) with CMI / Finance / Brand Category Tabs",
        "layer": "FE",
        "owner": "Vishwachandra",
        "start": "19-Oct-2026",
        "end": "01-Nov-2026",
        "weeks": ["W2", "W3"]
    },
    {
        "epic": "Main Page & Parameters",
        "feature": "Parameter Selection Form UI",
        "story": "Parameter Selection Form: Region, Market, Presentation Template dropdowns and client validation",
        "layer": "FE",
        "owner": "Vishwachandra",
        "start": "26-Oct-2026",
        "end": "08-Nov-2026",
        "weeks": ["W3", "W4"]
    },
    {
        "epic": "Main Page & Parameters",
        "feature": "Recent FNOS Reports UI",
        "story": "Recent FNOS Reports List Cards (FNOS_26_Q1/Q2/Q3, View: Total Europe Cluster, Status Badges)",
        "layer": "FE",
        "owner": "Vishwachandra",
        "start": "26-Oct-2026",
        "end": "08-Nov-2026",
        "weeks": ["W3", "W4"]
    },
    {
        "epic": "Reports Library & Canvas",
        "feature": "Accordion Reports Library UI",
        "story": "Reports Library Accordion Sidebar: EU Overview, Financial, Drive Relevance, SRM category tree navigation",
        "layer": "FE",
        "owner": "Sudhakshina",
        "start": "26-Oct-2026",
        "end": "08-Nov-2026",
        "weeks": ["W3", "W4"]
    },
    {
        "epic": "Reports Library & Canvas",
        "feature": "16:9 Matrix Canvas UI",
        "story": "Brand Performance Matrix 16:9 Canvas: Multi-tier headers (Seed, Growth, Fuel, Harvest) and YTD/MAT dual rows",
        "layer": "FE",
        "owner": "Sudhakshina",
        "start": "02-Nov-2026",
        "end": "08-Nov-2026",
        "weeks": ["W4"]
    },
    {
        "epic": "Distribution & PPTX Engine",
        "feature": "PPTX Template Decomposition",
        "story": "Deconstruction of Client 58-Slide Master Deck into Python-PPTX Slide Layouts and Native Table Stubs",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "02-Nov-2026",
        "end": "08-Nov-2026",
        "weeks": ["W4"]
    },

    # --- PHASE 2: Live Data Integration & Orchestration (Weeks 5 - 8: Nov 9 - Dec 6) ---
    {
        "epic": "Main Page & Parameters",
        "feature": "Database Connection & ORM",
        "story": "Azure SQL DB Connection and SQLAlchemy ORM Layer (Consuming DS-managed schema and application tables)",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "09-Nov-2026",
        "end": "22-Nov-2026",
        "weeks": ["W5", "W6"]
    },
    {
        "epic": "Main Page & Parameters",
        "feature": "Parameter & Report APIs",
        "story": "REST Endpoints for KBI metrics, filter options and recent report history (Consuming live DS tables)",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "09-Nov-2026",
        "end": "22-Nov-2026",
        "weeks": ["W5", "W6"]
    },
    {
        "epic": "Reports Library & Canvas",
        "feature": "Databricks Connector & APIs",
        "story": "Databricks SQL Warehouse Connector and REST Data Integration Service for Matrix Rollups",
        "layer": "BE",
        "owner": "Sudhakshina",
        "start": "09-Nov-2026",
        "end": "22-Nov-2026",
        "weeks": ["W5", "W6"]
    },
    {
        "epic": "AI Co-Pilot & Commentary",
        "feature": "Slide-by-Slide Text Editor",
        "story": "In-Canvas Slide-by-Slide Narrative Editing: Headline, So What, Now What and Save-to-Slide UI",
        "layer": "FE",
        "owner": "Sudhakshina",
        "start": "16-Nov-2026",
        "end": "29-Nov-2026",
        "weeks": ["W6", "W7"]
    },
    {
        "epic": "Reports Library & Canvas",
        "feature": "Query Optimization & Caching",
        "story": "Analytical Query In-Memory Caching, Fallback Resilience and Slide Payload Serialization",
        "layer": "BE",
        "owner": "Sudhakshina",
        "start": "16-Nov-2026",
        "end": "29-Nov-2026",
        "weeks": ["W6", "W7"]
    },
    {
        "epic": "AI Co-Pilot & Commentary",
        "feature": "Presentation Orchestrator Engine",
        "story": "Job Lifecycle State Machine (PENDING, RUNNING, COMPLETED) and Async Task Dispatcher",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "16-Nov-2026",
        "end": "29-Nov-2026",
        "weeks": ["W6", "W7"]
    },
    {
        "epic": "AI Co-Pilot & Commentary",
        "feature": "Real-Time Progress Streaming",
        "story": "Report Generation Job Progress Tracking via Server-Sent Events (SSE)",
        "layer": "FE + BE",
        "owner": "Vishwachandra",
        "start": "23-Nov-2026",
        "end": "06-Dec-2026",
        "weeks": ["W7", "W8"]
    },
    {
        "epic": "AI Co-Pilot & Commentary",
        "feature": "Right Side AI Drawer UI",
        "story": "Right Side AI Co-Pilot Drawer (Triggered by AI CO-PILOT): Prompt input and suggested prompt chips",
        "layer": "FE",
        "owner": "Vishwachandra",
        "start": "23-Nov-2026",
        "end": "06-Dec-2026",
        "weeks": ["W7", "W8"]
    },
    {
        "epic": "Reports Library & Canvas",
        "feature": "Multi-Report Template Sync",
        "story": "Interactive Slide Switching across 4 Report Folders (Matching client 58-slide master deck)",
        "layer": "FE + BE",
        "owner": "Sudhakshina",
        "start": "23-Nov-2026",
        "end": "06-Dec-2026",
        "weeks": ["W7", "W8"]
    },
    {
        "epic": "AI Co-Pilot & Commentary",
        "feature": "Slide Edit Persistence API",
        "story": "Slide Narrative Delta Persistence API (Storing user edits in DS-provided Azure SQL tables)",
        "layer": "BE",
        "owner": "Sudhakshina",
        "start": "30-Nov-2026",
        "end": "06-Dec-2026",
        "weeks": ["W8"]
    },

    # --- PHASE 3: GenAI Synthesis, PPTX Engine & Distribution (Weeks 9 - 12: Dec 7 - Jan 3) ---
    {
        "epic": "Distribution & PPTX Engine",
        "feature": "Distribution Modal Dialog UI",
        "story": "Distribute Report Modal Dialog: Title input, recipient group pills (Executive Board/Managers) and Summary",
        "layer": "FE",
        "owner": "Vishwachandra",
        "start": "07-Dec-2026",
        "end": "13-Dec-2026",
        "weeks": ["W9"]
    },
    {
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Azure OpenAI Client Integration",
        "story": "FastAPI Client for Azure OpenAI / AI Foundry with Pydantic Structured Outputs",
        "layer": "BE",
        "owner": "Sudhakshina",
        "start": "07-Dec-2026",
        "end": "20-Dec-2026",
        "weeks": ["W9", "W10"]
    },
    {
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Commentary Regeneration Pipeline",
        "story": "Commentary Synthesis API and Guardrail Enforcement (Integrating Jesu tuned prompts)",
        "layer": "BE",
        "owner": "Sudhakshina",
        "start": "07-Dec-2026",
        "end": "20-Dec-2026",
        "weeks": ["W9", "W10"]
    },
    {
        "epic": "Distribution & PPTX Engine",
        "feature": "AI Executive Summary Regeneration",
        "story": "AI Executive Summary Narrative Generation and Regenerate Action in Distribute Modal",
        "layer": "FE + BE",
        "owner": "Vishwachandra",
        "start": "07-Dec-2026",
        "end": "20-Dec-2026",
        "weeks": ["W9", "W10"]
    },
    {
        "epic": "Distribution & PPTX Engine",
        "feature": "Enterprise Email Service",
        "story": "Enterprise Email Integration (Microsoft Graph API / SendGrid): Dispatch Report to Selected Groups",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "14-Dec-2026",
        "end": "27-Dec-2026",
        "weeks": ["W10", "W11"]
    },
    {
        "epic": "Distribution & PPTX Engine",
        "feature": "Serverless 80-Slide PPTX Engine",
        "story": "Serverless PPTX Rendering Worker: Hydrating master deck templates across 4 report sections",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "14-Dec-2026",
        "end": "27-Dec-2026",
        "weeks": ["W10", "W11"]
    },
    {
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Conversational Insights Stream",
        "story": "Interactive Scenario Q&A API and Frontend Narrative Generation Controls",
        "layer": "FE + BE",
        "owner": "Sudhakshina",
        "start": "14-Dec-2026",
        "end": "27-Dec-2026",
        "weeks": ["W10", "W11"]
    },
    {
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Dynamic Chart Visualizations",
        "story": "Chart Components for Brand Share, Distribution and Spend Trends",
        "layer": "FE",
        "owner": "Sudhakshina",
        "start": "21-Dec-2026",
        "end": "03-Jan-2027",
        "weeks": ["W11", "W12"]
    },
    {
        "epic": "Distribution & PPTX Engine",
        "feature": "Blob SAS Download URL API",
        "story": "Azure Blob Storage Deck Upload and Presigned SAS Download Token API (One-click PPTX Download)",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "21-Dec-2026",
        "end": "03-Jan-2027",
        "weeks": ["W11", "W12"]
    },
    {
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Accessibility & Polish",
        "story": "Accessibility, Cross-Browser Compatibility and UI Polish",
        "layer": "FE",
        "owner": "Sudhakshina",
        "start": "28-Dec-2026",
        "end": "03-Jan-2027",
        "weeks": ["W12"]
    },

    # --- PHASE 4: Hardening, Latency Profiling & Final Code Freeze (Weeks 13 - 14: Jan 4 - Jan 17) ---
    {
        "epic": "Hardening & Code Freeze",
        "feature": "Frontend State Optimization",
        "story": "Frontend Rendering Optimization, Memory Leak Prevention, Error Boundaries and Toast System",
        "layer": "FE",
        "owner": "Sudhakshina",
        "start": "04-Jan-2027",
        "end": "10-Jan-2027",
        "weeks": ["W13"]
    },
    {
        "epic": "Hardening & Code Freeze",
        "feature": "Backend Query Profiling & Tuning",
        "story": "Analytical Query Profiling, In-Memory Caching Rules, Connection Pooling and Async Thread Tuning",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "04-Jan-2027",
        "end": "10-Jan-2027",
        "weeks": ["W13"]
    },
    {
        "epic": "Hardening & Code Freeze",
        "feature": "80-Slide Rendering Latency Tuning",
        "story": "80-Slide PowerPoint Generation Optimization (Sub-second target) and Memory Profiling",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "04-Jan-2027",
        "end": "17-Jan-2027",
        "weeks": ["W13", "W14"]
    },
    {
        "epic": "Hardening & Code Freeze",
        "feature": "Final Feature Freeze & Release Build",
        "story": "Final App Development Feature Freeze, Production Release Build and Handoff to Testing Team",
        "layer": "FE + BE",
        "owner": "Sudhakshina",
        "start": "11-Jan-2027",
        "end": "17-Jan-2027",
        "weeks": ["W14"]
    }
]

week_cols = [
    ("W1", "W1 (12-Oct)"),
    ("W2", "W2 (19-Oct)"),
    ("W3", "W3 (26-Oct)"),
    ("W4", "W4 (02-Nov) - M1"),
    ("W5", "W5 (09-Nov)"),
    ("W6", "W6 (16-Nov)"),
    ("W7", "W7 (23-Nov)"),
    ("W8", "W8 (30-Nov)"),
    ("W9", "W9 (07-Dec)"),
    ("W10", "W10 (14-Dec)"),
    ("W11", "W11 (21-Dec)"),
    ("W12", "W12 (28-Dec)"),
    ("W13", "W13 (04-Jan)"),
    ("W14", "W14 (11-Jan) - M2")
]

headers = ["Epic", "Feature", "User Story", "Layer", "Owner", "Start Date", "End Date"] + [w[1] for w in week_cols]

def populate_sheet_plain(ws):
    # Plain fonts: standard black font, regular weight for data, bold for headers
    font_header = Font(name="Calibri", size=10, bold=True, color="000000")
    font_data = Font(name="Calibri", size=10, bold=False, color="000000")
    font_gantt = Font(name="Calibri", size=10, bold=True, color="000000")

    # Plain fills: Header has light neutral gray F2F2F2, active weeks have light neutral gray EAEAEA, all else white/none
    fill_header = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    fill_active_week = PatternFill(start_color="EAEAEA", end_color="EAEAEA", fill_type="solid")

    thin_border_side = Side(style="thin", color="D9D9D9")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    header_border_side = Side(style="thin", color="BFBFBF")
    header_border = Border(left=header_border_side, right=header_border_side, top=header_border_side, bottom=header_border_side)

    # Append Header
    ws.append(headers)
    ws.row_dimensions[1].height = 24
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(1, col_idx)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = fill_header
        cell.font = font_header
        cell.border = header_border

    # Append Data Rows
    for r_idx, s in enumerate(stories, start=2):
        row_vals = [
            s["epic"],
            s["feature"],
            s["story"],
            s["layer"],
            s["owner"],
            s["start"],
            s["end"]
        ]
        for w_code, _ in week_cols:
            if w_code in s["weeks"]:
                row_vals.append("X")
            else:
                row_vals.append("")
        
        ws.append(row_vals)
        ws.row_dimensions[r_idx].height = 20
        
        # Meta columns (Cols 1 to 7) - strictly plain white, no colored fills, plain black text
        for c_idx in range(1, 8):
            cell = ws.cell(r_idx, c_idx)
            cell.font = font_data
            cell.border = thin_border
            if c_idx in (1, 2, 3):
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

        # Gantt week columns (Cols 8 to 21) - plain neutral gray fill with "X", or empty white
        for g_idx in range(8, len(headers) + 1):
            g_cell = ws.cell(r_idx, g_idx)
            g_cell.alignment = Alignment(horizontal="center", vertical="center")
            g_cell.border = thin_border
            if g_cell.value == "X":
                g_cell.fill = fill_active_week
                g_cell.font = font_gantt
            else:
                g_cell.font = font_data

    # Freeze panes at H2
    ws.freeze_panes = "H2"

    # Set column widths
    col_widths = {
        "A": 26, # Epic
        "B": 32, # Feature
        "C": 80, # User Story
        "D": 12, # Layer
        "E": 18, # Owner
        "F": 14, # Start Date
        "G": 14, # End Date
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width
    for w_col_idx in range(8, len(headers) + 1):
        col_letter = get_column_letter(w_col_idx)
        ws.column_dimensions[col_letter].width = 17

if __name__ == "__main__":
    wb_standalone = openpyxl.Workbook()
    ws_standalone = wb_standalone.active
    ws_standalone.title = "App Development Plan"
    populate_sheet_plain(ws_standalone)

    # 1. Save to dedicated Plain workbook
    wb_standalone.save("FNOS_App_Development_Plan_14Weeks_Plain.xlsx")
    print("Saved successfully to FNOS_App_Development_Plan_14Weeks_Plain.xlsx")

    # 2. Update master kickoff draft FNOS_Project_Plan_Initial_Draft.xlsx
    try:
        wb_draft = openpyxl.load_workbook("FNOS_Project_Plan_Initial_Draft.xlsx")
        if "Webapp Plan" in wb_draft.sheetnames:
            del wb_draft["Webapp Plan"]
        ws_draft = wb_draft.create_sheet(title="Webapp Plan")
        populate_sheet_plain(ws_draft)
        wb_draft.save("FNOS_Project_Plan_Initial_Draft.xlsx")
        print("Saved successfully to FNOS_Project_Plan_Initial_Draft.xlsx in 'Webapp Plan' tab!")
    except Exception as e:
        print("Error updating FNOS_Project_Plan_Initial_Draft.xlsx:", e)

    # 3. Try saving to v2 if user closes it
    try:
        wb_standalone.save("FNOS_App_Development_Plan_14Weeks_v2.xlsx")
        print("Saved successfully to FNOS_App_Development_Plan_14Weeks_v2.xlsx")
    except Exception as e:
        print("Notice: FNOS_App_Development_Plan_14Weeks_v2.xlsx currently open in Excel.")

    # 4. Try saving to original file if user closes it
    try:
        wb_standalone.save("FNOS_App_Development_Plan_14Weeks.xlsx")
        print("Saved successfully to FNOS_App_Development_Plan_14Weeks.xlsx")
    except Exception as e:
        print("Notice: FNOS_App_Development_Plan_14Weeks.xlsx currently open in Excel.")
