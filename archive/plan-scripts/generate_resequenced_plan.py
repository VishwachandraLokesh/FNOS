import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

stories = [
    # --- PHASE 1: Zero-Dependency Foundations & Interactive Prototype (Weeks 1 - 4: Oct 12 - Nov 8) ---
    {
        "phase": "Phase 1: Zero-Dependency Foundations & Interactive Prototype",
        "epic": "Foundations & Auth",
        "feature": "Dev Environment Setup",
        "story": "Monorepo & App Scaffolding: React 19 + Vite Frontend, FastAPI Backend, Pydantic & Linters",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "12-Oct-2026",
        "end": "18-Oct-2026",
        "weeks": ["W1"]
    },
    {
        "phase": "Phase 1: Zero-Dependency Foundations & Interactive Prototype",
        "epic": "Foundations & Auth",
        "feature": "Design System & UI Tokens",
        "story": "Mars Design System Setup: UI tokens, typography, 16:9 canvas container shell & floating food theme",
        "layer": "FE",
        "owner": "Sudhakshina",
        "start": "12-Oct-2026",
        "end": "25-Oct-2026",
        "weeks": ["W1", "W2"]
    },
    {
        "phase": "Phase 1: Zero-Dependency Foundations & Interactive Prototype",
        "epic": "Foundations & Auth",
        "feature": "Login Page UI",
        "story": "Login Page UI: Welcome card, Mars-Federated-SSO badge, email input (effem.com) & Sign-In button",
        "layer": "FE",
        "owner": "Sudhakshina",
        "start": "12-Oct-2026",
        "end": "25-Oct-2026",
        "weeks": ["W1", "W2"]
    },
    {
        "phase": "Phase 1: Zero-Dependency Foundations & Interactive Prototype",
        "epic": "Foundations & Auth",
        "feature": "API Contracts & Stubs",
        "story": "OpenAPI & Pydantic Data Contract Specifications + Standalone Mock Fixture Generator for Local Dev",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "12-Oct-2026",
        "end": "25-Oct-2026",
        "weeks": ["W1", "W2"]
    },
    {
        "phase": "Phase 1: Zero-Dependency Foundations & Interactive Prototype",
        "epic": "Foundations & Auth",
        "feature": "Federated SSO & Auth API",
        "story": "Azure AD / Entra ID Federated SSO Integration: MSAL.js Auth Context (FE) & JWT Validation Middleware",
        "layer": "FE + BE",
        "owner": "Sudhakshina",
        "start": "19-Oct-2026",
        "end": "01-Nov-2026",
        "weeks": ["W2", "W3"]
    },
    {
        "phase": "Phase 1: Zero-Dependency Foundations & Interactive Prototype",
        "epic": "Main Page & Parameters",
        "feature": "KBI Header Bar & Tabs UI",
        "story": "Main Page KBI Header Bar (RSV, Share, Volume, Pen.) with CMI / Finance / Brand Category Tabs (Mock Data)",
        "layer": "FE",
        "owner": "Vishwachandra",
        "start": "19-Oct-2026",
        "end": "01-Nov-2026",
        "weeks": ["W2", "W3"]
    },
    {
        "phase": "Phase 1: Zero-Dependency Foundations & Interactive Prototype",
        "epic": "Main Page & Parameters",
        "feature": "Parameter Selection Form UI",
        "story": "Parameter Selection Form: Region, Market, Presentation Template dropdowns & client-side validation",
        "layer": "FE",
        "owner": "Vishwachandra",
        "start": "26-Oct-2026",
        "end": "08-Nov-2026",
        "weeks": ["W3", "W4"]
    },
    {
        "phase": "Phase 1: Zero-Dependency Foundations & Interactive Prototype",
        "epic": "Main Page & Parameters",
        "feature": "Recent FNOS Reports UI",
        "story": "Recent FNOS Reports List Cards (FNOS_26_Q1/Q2/Q3, View: Total Europe Cluster, Published/Draft) (Mock Data)",
        "layer": "FE",
        "owner": "Vishwachandra",
        "start": "26-Oct-2026",
        "end": "08-Nov-2026",
        "weeks": ["W3", "W4"]
    },
    {
        "phase": "Phase 1: Zero-Dependency Foundations & Interactive Prototype",
        "epic": "Reports Library & Canvas",
        "feature": "Accordion Reports Library UI",
        "story": "Reports Library Accordion Sidebar: EU Overview, Financial, Drive Relevance, SRM category tree navigation",
        "layer": "FE",
        "owner": "Sudhakshina",
        "start": "26-Oct-2026",
        "end": "08-Nov-2026",
        "weeks": ["W3", "W4"]
    },
    {
        "phase": "Phase 1: Zero-Dependency Foundations & Interactive Prototype",
        "epic": "Reports Library & Canvas",
        "feature": "16:9 Matrix Canvas UI",
        "story": "Brand Performance Matrix 16:9 Canvas: Multi-tier headers (Seed, Growth, Fuel, Harvest) & YTD/MAT dual rows",
        "layer": "FE",
        "owner": "Sudhakshina",
        "start": "02-Nov-2026",
        "end": "08-Nov-2026",
        "weeks": ["W4"]
    },
    {
        "phase": "Phase 1: Zero-Dependency Foundations & Interactive Prototype",
        "epic": "Distribution & PPTX Engine",
        "feature": "PPTX Template Decomposition",
        "story": "Deconstruction of Client 58-Slide Master Deck into Python-PPTX Slide Layouts & Native Table Stubs",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "02-Nov-2026",
        "end": "08-Nov-2026",
        "weeks": ["W4"]
    },

    # --- PHASE 2: Live Data Integration & Orchestration (Weeks 5 - 8: Nov 9 - Dec 6) ---
    {
        "phase": "Phase 2: Live Data Integration & Orchestration",
        "epic": "Main Page & Parameters",
        "feature": "Database Connection & ORM",
        "story": "Azure SQL DB Connection & SQLAlchemy ORM Layer (Consuming DS-managed schema & application tables)",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "09-Nov-2026",
        "end": "22-Nov-2026",
        "weeks": ["W5", "W6"]
    },
    {
        "phase": "Phase 2: Live Data Integration & Orchestration",
        "epic": "Main Page & Parameters",
        "feature": "Parameter & Report APIs",
        "story": "REST Endpoints for KBI metrics, filter options & recent report history (Consuming live DS tables)",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "09-Nov-2026",
        "end": "22-Nov-2026",
        "weeks": ["W5", "W6"]
    },
    {
        "phase": "Phase 2: Live Data Integration & Orchestration",
        "epic": "Reports Library & Canvas",
        "feature": "Databricks Connector & APIs",
        "story": "Databricks SQL Warehouse Connector & REST Data Integration Service for Matrix Rollups",
        "layer": "BE",
        "owner": "Sudhakshina",
        "start": "09-Nov-2026",
        "end": "22-Nov-2026",
        "weeks": ["W5", "W6"]
    },
    {
        "phase": "Phase 2: Live Data Integration & Orchestration",
        "epic": "AI Co-Pilot & Commentary",
        "feature": "Slide-by-Slide Text Editor",
        "story": "In-Canvas Slide-by-Slide Narrative Editing: Headline, So What, Now What & Save-to-Slide UI",
        "layer": "FE",
        "owner": "Sudhakshina",
        "start": "16-Nov-2026",
        "end": "29-Nov-2026",
        "weeks": ["W6", "W7"]
    },
    {
        "phase": "Phase 2: Live Data Integration & Orchestration",
        "epic": "Reports Library & Canvas",
        "feature": "Query Optimization & Caching",
        "story": "Analytical Query In-Memory Caching, Fallback Resilience & Slide Payload Serialization",
        "layer": "BE",
        "owner": "Sudhakshina",
        "start": "16-Nov-2026",
        "end": "29-Nov-2026",
        "weeks": ["W6", "W7"]
    },
    {
        "phase": "Phase 2: Live Data Integration & Orchestration",
        "epic": "AI Co-Pilot & Commentary",
        "feature": "Presentation Orchestrator Engine",
        "story": "Job Lifecycle State Machine (PENDING, RUNNING, COMPLETED) & Async Task Dispatcher",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "16-Nov-2026",
        "end": "29-Nov-2026",
        "weeks": ["W6", "W7"]
    },
    {
        "phase": "Phase 2: Live Data Integration & Orchestration",
        "epic": "AI Co-Pilot & Commentary",
        "feature": "Real-Time Progress Streaming",
        "story": "Real-Time Job Progress Streaming via Server-Sent Events (SSE) / WebSockets",
        "layer": "FE + BE",
        "owner": "Vishwachandra",
        "start": "23-Nov-2026",
        "end": "06-Dec-2026",
        "weeks": ["W7", "W8"]
    },
    {
        "phase": "Phase 2: Live Data Integration & Orchestration",
        "epic": "AI Co-Pilot & Commentary",
        "feature": "Right Slide-Out AI Drawer UI",
        "story": "Right Slide-Out AI Co-Pilot Drawer (Triggered by ✦ AI CO-PILOT): Prompt input & suggested chips",
        "layer": "FE",
        "owner": "Vishwachandra",
        "start": "23-Nov-2026",
        "end": "06-Dec-2026",
        "weeks": ["W7", "W8"]
    },
    {
        "phase": "Phase 2: Live Data Integration & Orchestration",
        "epic": "Reports Library & Canvas",
        "feature": "Multi-Report Template Sync",
        "story": "Interactive Slide Switching across 4 Report Folders (Matching client 58-slide master deck)",
        "layer": "FE + BE",
        "owner": "Sudhakshina",
        "start": "23-Nov-2026",
        "end": "06-Dec-2026",
        "weeks": ["W7", "W8"]
    },
    {
        "phase": "Phase 2: Live Data Integration & Orchestration",
        "epic": "AI Co-Pilot & Commentary",
        "feature": "Slide Edit Persistence API",
        "story": "Slide Narrative Delta Persistence API (Storing user edits in DS-provided Azure SQL tables)",
        "layer": "BE",
        "owner": "Sudhakshina",
        "start": "30-Nov-2026",
        "end": "06-Dec-2026",
        "weeks": ["W8"]
    },

    # --- PHASE 3: GenAI Synthesis, PPTX Engine & Distribution (Weeks 9 - 12: Dec 7 - Jan 3) ---
    {
        "phase": "Phase 3: GenAI Synthesis, PPTX Engine & Distribution",
        "epic": "Distribution & PPTX Engine",
        "feature": "Distribution Modal Dialog UI",
        "story": "Distribute Report Modal Dialog: Title input, group pills (Executive Board/Managers) & Summary",
        "layer": "FE",
        "owner": "Vishwachandra",
        "start": "07-Dec-2026",
        "end": "13-Dec-2026",
        "weeks": ["W9"]
    },
    {
        "phase": "Phase 3: GenAI Synthesis, PPTX Engine & Distribution",
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Azure OpenAI Client Integration",
        "story": "FastAPI Client for Azure OpenAI / AI Foundry with Pydantic Structured Outputs",
        "layer": "BE",
        "owner": "Sudhakshina",
        "start": "07-Dec-2026",
        "end": "20-Dec-2026",
        "weeks": ["W9", "W10"]
    },
    {
        "phase": "Phase 3: GenAI Synthesis, PPTX Engine & Distribution",
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Commentary Regeneration Pipeline",
        "story": "Commentary Synthesis API & Guardrail Enforcement (Integrating Jesu tuned prompts)",
        "layer": "BE",
        "owner": "Sudhakshina",
        "start": "07-Dec-2026",
        "end": "20-Dec-2026",
        "weeks": ["W9", "W10"]
    },
    {
        "phase": "Phase 3: GenAI Synthesis, PPTX Engine & Distribution",
        "epic": "Distribution & PPTX Engine",
        "feature": "AI Executive Summary Regeneration",
        "story": "AI Executive Summary Narrative Generation & Regenerate Action in Distribute Modal",
        "layer": "FE + BE",
        "owner": "Vishwachandra",
        "start": "07-Dec-2026",
        "end": "20-Dec-2026",
        "weeks": ["W9", "W10"]
    },
    {
        "phase": "Phase 3: GenAI Synthesis, PPTX Engine & Distribution",
        "epic": "Distribution & PPTX Engine",
        "feature": "Enterprise Email Service",
        "story": "Enterprise Email Integration (Microsoft Graph API / SendGrid): Dispatch Report to Selected Groups",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "14-Dec-2026",
        "end": "27-Dec-2026",
        "weeks": ["W10", "W11"]
    },
    {
        "phase": "Phase 3: GenAI Synthesis, PPTX Engine & Distribution",
        "epic": "Distribution & PPTX Engine",
        "feature": "Serverless 80-Slide PPTX Engine",
        "story": "Serverless PPTX Rendering Worker: Hydrating master deck templates across 4 report sections",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "14-Dec-2026",
        "end": "27-Dec-2026",
        "weeks": ["W10", "W11"]
    },
    {
        "phase": "Phase 3: GenAI Synthesis, PPTX Engine & Distribution",
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Conversational Insights Stream",
        "story": "Interactive Scenario Q&A API + Frontend Live Streaming Narrative Regeneration Controls",
        "layer": "FE + BE",
        "owner": "Sudhakshina",
        "start": "14-Dec-2026",
        "end": "27-Dec-2026",
        "weeks": ["W10", "W11"]
    },
    {
        "phase": "Phase 3: GenAI Synthesis, PPTX Engine & Distribution",
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Dynamic Chart Visualizations",
        "story": "In-Canvas Interactive Chart Components for Brand Share, Distribution & Spend Trends",
        "layer": "FE",
        "owner": "Sudhakshina",
        "start": "21-Dec-2026",
        "end": "03-Jan-2027",
        "weeks": ["W11", "W12"]
    },
    {
        "phase": "Phase 3: GenAI Synthesis, PPTX Engine & Distribution",
        "epic": "Distribution & PPTX Engine",
        "feature": "Blob SAS Download URL API",
        "story": "Azure Blob Storage Deck Upload & Presigned SAS Download Token API (One-click PPTX Download)",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "21-Dec-2026",
        "end": "03-Jan-2027",
        "weeks": ["W11", "W12"]
    },
    {
        "phase": "Phase 3: GenAI Synthesis, PPTX Engine & Distribution",
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Accessibility & Responsive Polish",
        "story": "Executive Ready Mode Toggle, WCAG 2.1 Accessibility & Cross-Browser Compatibility Polish",
        "layer": "FE",
        "owner": "Sudhakshina",
        "start": "28-Dec-2026",
        "end": "03-Jan-2027",
        "weeks": ["W12"]
    },

    # --- PHASE 4: Hardening, Latency Profiling & Final Code Freeze (Weeks 13 - 14: Jan 4 - Jan 17) ---
    {
        "phase": "Phase 4: Hardening, Latency Profiling & Final Code Freeze",
        "epic": "Hardening & Code Freeze",
        "feature": "Frontend State Optimization",
        "story": "Frontend Rendering Optimization, Memory Leak Prevention, Error Boundaries & Toast System",
        "layer": "FE",
        "owner": "Sudhakshina",
        "start": "04-Jan-2027",
        "end": "10-Jan-2027",
        "weeks": ["W13"]
    },
    {
        "phase": "Phase 4: Hardening, Latency Profiling & Final Code Freeze",
        "epic": "Hardening & Code Freeze",
        "feature": "Backend Query Profiling & Tuning",
        "story": "Analytical Query Profiling, In-Memory Caching Rules, Connection Pooling & Async Thread Tuning",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "04-Jan-2027",
        "end": "10-Jan-2027",
        "weeks": ["W13"]
    },
    {
        "phase": "Phase 4: Hardening, Latency Profiling & Final Code Freeze",
        "epic": "Hardening & Code Freeze",
        "feature": "80-Slide Rendering Latency Tuning",
        "story": "80-Slide PowerPoint Generation Optimization (Sub-second target) & Memory Profiling",
        "layer": "BE",
        "owner": "Vishwachandra",
        "start": "04-Jan-2027",
        "end": "17-Jan-2027",
        "weeks": ["W13", "W14"]
    },
    {
        "phase": "Phase 4: Hardening, Latency Profiling & Final Code Freeze",
        "epic": "Hardening & Code Freeze",
        "feature": "Final Feature Freeze & Release Build",
        "story": "Final App Development Feature Freeze, Production Release Build & Handoff to Testing Team",
        "layer": "FE + BE",
        "owner": "Sudhakshina",
        "start": "11-Jan-2027",
        "end": "17-Jan-2027",
        "weeks": ["W14"]
    }
]

week_cols = [
    ("W1", "W1 (12-Oct)"),
    ("W2", "W2 (19-Oct)"),
    ("W3", "W3 (26-Oct)"),
    ("W4", "W4 (02-Nov) - M1"),
    ("W5", "W5 (09-Nov)"),
    ("W6", "W6 (16-Nov)"),
    ("W7", "W7 (23-Nov)"),
    ("W8", "W8 (30-Nov)"),
    ("W9", "W9 (07-Dec)"),
    ("W10", "W10 (14-Dec)"),
    ("W11", "W11 (21-Dec)"),
    ("W12", "W12 (28-Dec)"),
    ("W13", "W13 (04-Jan)"),
    ("W14", "W14 (11-Jan) - M2")
]

headers = ["Epic", "Feature", "User Story", "Layer", "Owner", "Start Date", "End Date"] + [w[1] for w in week_cols]

def populate_sheet(ws):
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10)
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_owner_v = Font(name="Calibri", size=10, color="002060", bold=True)
    font_owner_s = Font(name="Calibri", size=10, color="5B2C6F", bold=True)
    
    font_layer_fe = Font(name="Calibri", size=10, bold=True, color="004085")
    fill_layer_fe = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")
    
    font_layer_be = Font(name="Calibri", size=10, bold=True, color="155724")
    fill_layer_be = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    
    font_layer_fs = Font(name="Calibri", size=10, bold=True, color="856404")
    fill_layer_fs = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")

    fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    fill_header_m1 = PatternFill(start_color="27AE60", end_color="27AE60", fill_type="solid")
    fill_header_m2 = PatternFill(start_color="C0392B", end_color="C0392B", fill_type="solid")
    
    fill_gantt_v = PatternFill(start_color="B4C6E7", end_color="B4C6E7", fill_type="solid")
    font_gantt_v = Font(name="Calibri", size=10, bold=True, color="1F497D")
    
    fill_mint = PatternFill(start_color="D9F2D0", end_color="D9F2D0", fill_type="solid")
    font_mint = Font(name="Calibri", size=10, bold=True, color="274E13")

    thin_border_side = Side(style="thin", color="D9D9D9")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    header_border_side = Side(style="medium", color="1F497D")
    header_border = Border(left=thin_border_side, right=thin_border_side, top=header_border_side, bottom=header_border_side)

    # Append Header
    ws.append(headers)
    ws.row_dimensions[1].height = 28
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(1, col_idx)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        w_label = headers[col_idx - 1]
        if "M1" in w_label:
            cell.fill = fill_header_m1
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        elif "M2" in w_label:
            cell.fill = fill_header_m2
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        else:
            cell.fill = fill_header
            cell.font = font_header
        cell.border = header_border

    # Append Data Rows
    for r_idx, s in enumerate(stories, start=2):
        row_vals = [
            s["epic"],
            s["feature"],
            s["story"],
            s["layer"],
            s["owner"],
            s["start"],
            s["end"]
        ]
        for w_code, _ in week_cols:
            if w_code in s["weeks"]:
                row_vals.append("✓")
            else:
                row_vals.append("")
        
        ws.append(row_vals)
        ws.row_dimensions[r_idx].height = 22
        
        c1 = ws.cell(r_idx, 1)
        c2 = ws.cell(r_idx, 2)
        c3 = ws.cell(r_idx, 3)
        c4 = ws.cell(r_idx, 4)
        c5 = ws.cell(r_idx, 5)
        c6 = ws.cell(r_idx, 6)
        c7 = ws.cell(r_idx, 7)

        c1.alignment = Alignment(horizontal="left", vertical="center")
        c2.alignment = Alignment(horizontal="left", vertical="center")
        c3.alignment = Alignment(horizontal="left", vertical="center")
        c4.alignment = Alignment(horizontal="center", vertical="center")
        c5.alignment = Alignment(horizontal="center", vertical="center")
        c6.alignment = Alignment(horizontal="center", vertical="center")
        c7.alignment = Alignment(horizontal="center", vertical="center")

        c1.font = font_data
        c2.font = font_data
        c3.font = font_bold
        
        if s["layer"] == "FE":
            c4.fill = fill_layer_fe
            c4.font = font_layer_fe
        elif s["layer"] == "BE":
            c4.fill = fill_layer_be
            c4.font = font_layer_be
        else:
            c4.fill = fill_layer_fs
            c4.font = font_layer_fs

        if s["owner"] == "Vishwachandra":
            c5.font = font_owner_v
        else:
            c5.font = font_owner_s

        c6.font = font_data
        c7.font = font_data

        for c_idx in range(1, 8):
            ws.cell(r_idx, c_idx).border = thin_border

        for g_idx in range(8, len(headers) + 1):
            g_cell = ws.cell(r_idx, g_idx)
            g_cell.alignment = Alignment(horizontal="center", vertical="center")
            g_cell.border = thin_border
            if g_cell.value == "✓":
                if s["owner"] == "Vishwachandra":
                    g_cell.fill = fill_gantt_v
                    g_cell.font = font_gantt_v
                else:
                    g_cell.fill = fill_mint
                    g_cell.font = font_mint

    ws.freeze_panes = "H2"

    col_widths = {
        "A": 26,
        "B": 32,
        "C": 80,
        "D": 14,
        "E": 18,
        "F": 14,
        "G": 14,
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width
    for w_col_idx in range(8, len(headers) + 1):
        col_letter = get_column_letter(w_col_idx)
        ws.column_dimensions[col_letter].width = 17

if __name__ == "__main__":
    wb_standalone = openpyxl.Workbook()
    ws_standalone = wb_standalone.active
    ws_standalone.title = "App Development Plan"
    populate_sheet(ws_standalone)

    wb_standalone.save("FNOS_App_Development_Plan_14Weeks_v2.xlsx")
    print("Saved successfully to FNOS_App_Development_Plan_14Weeks_v2.xlsx")

    try:
        wb_standalone.save("FNOS_App_Development_Plan_14Weeks.xlsx")
        print("Saved successfully to FNOS_App_Development_Plan_14Weeks.xlsx")
    except Exception as e:
        print("FNOS_App_Development_Plan_14Weeks.xlsx lock notice:", e)

    try:
        wb_draft = openpyxl.load_workbook("FNOS_Project_Plan_Initial_Draft.xlsx")
        if "Webapp Plan" in wb_draft.sheetnames:
            del wb_draft["Webapp Plan"]
        ws_draft = wb_draft.create_sheet(title="Webapp Plan")
        populate_sheet(ws_draft)
        wb_draft.save("FNOS_Project_Plan_Initial_Draft.xlsx")
        print("Saved successfully to FNOS_Project_Plan_Initial_Draft.xlsx in 'Webapp Plan' tab!")
    except Exception as e:
        print("Error updating FNOS_Project_Plan_Initial_Draft.xlsx:", e)
