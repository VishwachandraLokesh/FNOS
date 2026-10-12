import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import copy

stories = [
    # --- PHASE 1: Zero-Dependency Foundations & Interactive Prototype (Weeks 1 - 4: Oct 12 - Nov 8) ---
    {
        "epic": "Foundations & Auth",
        "feature": "Dev Environment Setup",
        "story": "Monorepo and App Scaffolding: React 19 + Vite Frontend, FastAPI Backend, Pydantic and Linters",
        "layer": "BE",
        "owner": "Vishwachandra",
        "weeks": ["W1"]
    },
    {
        "epic": "Foundations & Auth",
        "feature": "UI Framework & Base Layout",
        "story": "Setup Base Application Styling, Common UI Components and Page Layout",
        "layer": "FE",
        "owner": "Sudhakshina",
        "weeks": ["W1", "W2"]
    },
    {
        "epic": "Foundations & Auth",
        "feature": "Login Page UI",
        "story": "Login Page UI: Welcome card, Mars-Federated-SSO badge, email input (effem.com) and Sign-In button",
        "layer": "FE",
        "owner": "Sudhakshina",
        "weeks": ["W1", "W2"]
    },
    {
        "epic": "Foundations & Auth",
        "feature": "API Contracts & Mock Services",
        "story": "Define API Data Models and Setup Mock APIs for Frontend Development",
        "layer": "BE",
        "owner": "Vishwachandra",
        "weeks": ["W1", "W2"]
    },
    {
        "epic": "Foundations & Auth",
        "feature": "Federated SSO & Auth API",
        "story": "Azure AD / Entra ID Federated SSO Integration: MSAL.js Auth Context (FE) and JWT Validation Middleware",
        "layer": "FE + BE",
        "owner": "Sudhakshina",
        "weeks": ["W2", "W3"]
    },
    {
        "epic": "Main Page & Parameters",
        "feature": "KBI Header Bar & Tabs UI",
        "story": "Main Page KBI Header Bar (RSV, Share, Volume, Pen.) with CMI / Finance / Brand Category Tabs",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W2", "W3"]
    },
    {
        "epic": "Main Page & Parameters",
        "feature": "Parameter Selection Form UI",
        "story": "Parameter Selection Form: Region, Market, Presentation Template dropdowns and client validation",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W3", "W4"]
    },
    {
        "epic": "Main Page & Parameters",
        "feature": "Recent FNOS Reports UI",
        "story": "Recent FNOS Reports List Cards (FNOS_26_Q1/Q2/Q3, View: Total Europe Cluster, Status Badges)",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W3", "W4"]
    },
    {
        "epic": "Presentation Canvas & Slides",
        "feature": "Slide Thumbnail Navigation",
        "story": "Left Sidebar Slide Thumbnail Navigation: Mini slide preview list and active slide selection",
        "layer": "FE",
        "owner": "Sudhakshina",
        "weeks": ["W3", "W4"]
    },
    {
        "epic": "Presentation Canvas & Slides",
        "feature": "16:9 Slide Canvas Viewer",
        "story": "Main Slide Canvas: 16:9 Slide Viewer and Interactive Data Display",
        "layer": "FE",
        "owner": "Sudhakshina",
        "weeks": ["W4"]
    },
    {
        "epic": "Distribution & PPTX Engine",
        "feature": "PowerPoint Template Setup",
        "story": "PowerPoint Template Engine: Deconstruct Master Deck into Reusable Slide Layouts",
        "layer": "BE",
        "owner": "Vishwachandra",
        "weeks": ["W4"]
    },

    # --- PHASE 2: Live Data Integration & Orchestration (Weeks 5 - 8: Nov 9 - Dec 6) ---
    {
        "epic": "Main Page & Parameters",
        "feature": "Database Connection & ORM",
        "story": "Azure SQL DB Connection and SQLAlchemy ORM Layer (Consuming DS-managed schema and application tables)",
        "layer": "BE",
        "owner": "Vishwachandra",
        "weeks": ["W5", "W6"]
    },
    {
        "epic": "Main Page & Parameters",
        "feature": "Parameter & Report APIs",
        "story": "REST Endpoints for KBI metrics, filter options and recent report history (Consuming live DS tables)",
        "layer": "BE",
        "owner": "Vishwachandra",
        "weeks": ["W5", "W6"]
    },
    {
        "epic": "Presentation Canvas & Slides",
        "feature": "Databricks Connector & APIs",
        "story": "Databricks SQL Warehouse Connector and REST Data Integration Service for Matrix Rollups",
        "layer": "BE",
        "owner": "Sudhakshina",
        "weeks": ["W5", "W6"]
    },
    {
        "epic": "AI Co-Pilot & Commentary",
        "feature": "Slide Commentary Editor",
        "story": "Slide Commentary Editor: In-Canvas Narrative Editing and Save-to-Slide UI",
        "layer": "FE",
        "owner": "Sudhakshina",
        "weeks": ["W6", "W7"]
    },
    {
        "epic": "Presentation Canvas & Slides",
        "feature": "Query Optimization & Caching",
        "story": "Analytical Query In-Memory Caching, Fallback Resilience and Slide Payload Serialization",
        "layer": "BE",
        "owner": "Sudhakshina",
        "weeks": ["W6", "W7"]
    },
    {
        "epic": "Presentation Canvas & Slides",
        "feature": "Presentation Orchestrator Engine",
        "story": "Job Lifecycle State Machine (PENDING, RUNNING, COMPLETED) and Async Task Dispatcher",
        "layer": "BE",
        "owner": "Vishwachandra",
        "weeks": ["W6", "W7"]
    },
    {
        "epic": "Presentation Canvas & Slides",
        "feature": "Real-Time Progress Streaming",
        "story": "Report Generation Job Progress Tracking via Server-Sent Events (SSE)",
        "layer": "FE + BE",
        "owner": "Vishwachandra",
        "weeks": ["W7", "W8"]
    },
    {
        "epic": "AI Co-Pilot & Commentary",
        "feature": "Right Side AI & Commentary Panel",
        "story": "Right Side Panel: AI Insights Assistant and Slide Commentary Editor",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W7", "W8"]
    },
    {
        "epic": "Presentation Canvas & Slides",
        "feature": "Deck Slide Synchronization",
        "story": "Deck Navigation: Active Slide Switching and Data Synchronization across 80 Slides",
        "layer": "FE + BE",
        "owner": "Sudhakshina",
        "weeks": ["W7", "W8"]
    },
    {
        "epic": "AI Co-Pilot & Commentary",
        "feature": "Slide Edit Persistence API",
        "story": "Slide Narrative Delta Persistence API (Storing user edits in DS-provided Azure SQL tables)",
        "layer": "BE",
        "owner": "Sudhakshina",
        "weeks": ["W8"]
    },

    # --- PHASE 3: GenAI Synthesis, PPTX Engine & Distribution (Weeks 9 - 12: Dec 7 - Jan 3) ---
    {
        "epic": "Distribution & PPTX Engine",
        "feature": "Distribution Modal Dialog UI",
        "story": "Distribute Report Modal Dialog: Title input, recipient group pills (Executive Board/Managers) and Summary",
        "layer": "FE",
        "owner": "Vishwachandra",
        "weeks": ["W9"]
    },
    {
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Azure OpenAI Client Integration",
        "story": "FastAPI Client for Azure OpenAI / AI Foundry with Pydantic Structured Outputs",
        "layer": "BE",
        "owner": "Sudhakshina",
        "weeks": ["W9", "W10"]
    },
    {
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Commentary Regeneration Pipeline",
        "story": "Commentary Synthesis API and Guardrail Enforcement (Integrating Jesu tuned prompts)",
        "layer": "BE",
        "owner": "Sudhakshina",
        "weeks": ["W9", "W10"]
    },
    {
        "epic": "Distribution & PPTX Engine",
        "feature": "AI Executive Summary Regeneration",
        "story": "AI Executive Summary Narrative Generation and Regenerate Action in Distribute Modal",
        "layer": "FE + BE",
        "owner": "Vishwachandra",
        "weeks": ["W9", "W10"]
    },
    {
        "epic": "Distribution & PPTX Engine",
        "feature": "Enterprise Email Service",
        "story": "Enterprise Email Integration (Microsoft Graph API / SendGrid): Dispatch Report to Selected Groups",
        "layer": "BE",
        "owner": "Vishwachandra",
        "weeks": ["W10", "W11"]
    },
    {
        "epic": "Distribution & PPTX Engine",
        "feature": "Serverless 80-Slide PPTX Engine",
        "story": "Serverless PPTX Rendering Worker: Hydrating 80-Slide Master Deck with Processed Analytics",
        "layer": "BE",
        "owner": "Vishwachandra",
        "weeks": ["W10", "W11"]
    },
    {
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Conversational Insights Stream",
        "story": "Interactive Scenario Q&A API and Frontend Narrative Generation Controls",
        "layer": "FE + BE",
        "owner": "Sudhakshina",
        "weeks": ["W10", "W11"]
    },
    {
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Dynamic Chart Visualizations",
        "story": "Chart Components for Brand Share, Distribution and Spend Trends",
        "layer": "FE",
        "owner": "Sudhakshina",
        "weeks": ["W11", "W12"]
    },
    {
        "epic": "Distribution & PPTX Engine",
        "feature": "Blob SAS Download URL API",
        "story": "Azure Blob Storage Deck Upload and Presigned SAS Download Token API (One-click PPTX Download)",
        "layer": "BE",
        "owner": "Vishwachandra",
        "weeks": ["W11", "W12"]
    },
    {
        "epic": "GenAI Synthesis & Copilot",
        "feature": "Accessibility & Polish",
        "story": "Accessibility, Cross-Browser Compatibility and UI Polish",
        "layer": "FE",
        "owner": "Sudhakshina",
        "weeks": ["W12"]
    },

    # --- PHASE 4: Hardening, Latency Profiling & Final Code Freeze (Weeks 13 - 14: Jan 4 - Jan 17) ---
    {
        "epic": "Hardening & Code Freeze",
        "feature": "Frontend State Optimization",
        "story": "Frontend Rendering Optimization, Memory Leak Prevention, Error Boundaries and Toast System",
        "layer": "FE",
        "owner": "Sudhakshina",
        "weeks": ["W13"]
    },
    {
        "epic": "Hardening & Code Freeze",
        "feature": "Backend Query Profiling & Tuning",
        "story": "Analytical Query Profiling, In-Memory Caching Rules, Connection Pooling and Async Thread Tuning",
        "layer": "BE",
        "owner": "Vishwachandra",
        "weeks": ["W13"]
    },
    {
        "epic": "Hardening & Code Freeze",
        "feature": "80-Slide Rendering Latency Tuning",
        "story": "80-Slide PowerPoint Generation Optimization (Sub-second target) and Memory Profiling",
        "layer": "BE",
        "owner": "Vishwachandra",
        "weeks": ["W13", "W14"]
    },
    {
        "epic": "Hardening & Code Freeze",
        "feature": "Final Feature Freeze & Release Build",
        "story": "Final App Development Feature Freeze, Production Release Build and Handoff to Testing Team",
        "layer": "FE + BE",
        "owner": "Sudhakshina",
        "weeks": ["W14"]
    }
]

week_cols = [
    ("W1", "W1"),
    ("W2", "W2"),
    ("W3", "W3"),
    ("W4", "W4 - M1"),
    ("W5", "W5"),
    ("W6", "W6"),
    ("W7", "W7"),
    ("W8", "W8"),
    ("W9", "W9"),
    ("W10", "W10"),
    ("W11", "W11"),
    ("W12", "W12"),
    ("W13", "W13"),
    ("W14", "W14")
]

headers = ["Epic", "Feature", "User Story", "Layer", "Owner"] + [w[1] for w in week_cols]

def build_sheet(ws, ws_ref):
    font_header = copy.copy(ws_ref.cell(1, 1).font)
    font_data = copy.copy(ws_ref.cell(2, 1).font)
    border_cell = copy.copy(ws_ref.cell(1, 1).border)
    fill_empty = PatternFill(fill_type=None)
    fill_m1 = copy.copy(ws_ref.cell(1, 8).fill)
    fill_active = copy.copy(ws_ref.cell(2, 5).fill)

    ws.append(headers)
    ws.row_dimensions[1].height = 14.25

    for c_idx in range(1, len(headers) + 1):
        cell = ws.cell(1, c_idx)
        cell.font = font_header
        cell.border = border_cell
        if headers[c_idx - 1] == "W4 - M1":
            cell.fill = fill_m1
        else:
            cell.fill = fill_empty
        if c_idx >= 6:
            cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx, s in enumerate(stories, start=2):
        row_vals = [
            s["epic"],
            s["feature"],
            s["story"],
            s["layer"],
            s["owner"]
        ]
        for _ in week_cols:
            row_vals.append(None)
            
        ws.append(row_vals)
        ws.row_dimensions[r_idx].height = 14.25

        for c_idx in range(1, 6):
            cell = ws.cell(r_idx, c_idx)
            cell.font = font_data
            cell.border = border_cell
            cell.fill = fill_empty
            if c_idx in (2, 3):
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            elif c_idx in (4, 5):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(vertical="center")

        for w_offset, (w_code, _) in enumerate(week_cols):
            col_pos = 6 + w_offset
            w_cell = ws.cell(r_idx, col_pos)
            w_cell.border = border_cell
            if w_code in s["weeks"]:
                w_cell.fill = fill_active
            else:
                w_cell.fill = fill_empty

    col_widths = {
        "A": 22.0,
        "B": 28.38,
        "C": 48.0,
        "D": 10.0,
        "E": 20.38,
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width
    for w_col_idx in range(6, len(headers) + 1):
        col_letter = get_column_letter(w_col_idx)
        if headers[w_col_idx - 1] == "W4 - M1":
            ws.column_dimensions[col_letter].width = 6.13
        else:
            ws.column_dimensions[col_letter].width = 5.0

if __name__ == "__main__":
    wb_master = openpyxl.load_workbook("FNOS_Project_Plan_Initial_Draft.xlsx")
    ws_ds = wb_master["DS Team Plan"]

    if "Webapp Plan" in wb_master.sheetnames:
        del wb_master["Webapp Plan"]
    ws_app = wb_master.create_sheet(title="Webapp Plan")
    build_sheet(ws_app, ws_ds)

    wb_master.save("FNOS_Project_Plan_Initial_Draft.xlsx")
    print("Updated FNOS_Project_Plan_Initial_Draft.xlsx [Webapp Plan] successfully!")

    wb_standalone = openpyxl.Workbook()
    ws_standalone = wb_standalone.active
    ws_standalone.title = "Webapp Plan"
    build_sheet(ws_standalone, ws_ds)
    wb_standalone.save("FNOS_App_Development_Plan_Initial_Draft_Style.xlsx")
    print("Created FNOS_App_Development_Plan_Initial_Draft_Style.xlsx")

    for tgt in ["FNOS_App_Development_Plan_14Weeks.xlsx", "FNOS_App_Development_Plan_14Weeks_v2.xlsx"]:
        try:
            wb_standalone.save(tgt)
            print(f"Updated {tgt}")
        except Exception:
            pass
